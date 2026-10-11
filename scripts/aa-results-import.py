#!/usr/bin/env python3
"""A&A Lote 7 — ingestão técnica reexecutável das folhas de respostas do Simula SAEB (adaptador simula-saeb-folha-respostas v1).
Modo padrão = DRY-RUN (só contagens agregadas, sem dado pessoal). Com --sql imprime SQL de staging idempotente:
  python3 scripts/aa-results-import.py --sql --program simula-saeb --year 2026 --round 3 > /tmp/aa.sql && psql -v ON_ERROR_STOP=1 -f /tmp/aa.sql
Regras: nome nunca vincula sozinho (vai à revisão humana); INEP só casa se presente; acertos locais nunca viram proficiência TRI;
a fórmula própria do instrumento (célula de total) é preservada como texto. Nenhum resultado canônico é gravado aqui."""
import argparse, glob, hashlib, json, re, subprocess, sys, unicodedata
import openpyxl

AD, VER, EX = "simula-saeb-folha-respostas", 1, "importacao-tecnica:scripts/aa-results-import.py"
ap = argparse.ArgumentParser(); ap.add_argument("--sql", action="store_true"); ap.add_argument("--program", default="simula-saeb")
ap.add_argument("--year", type=int, required=True); ap.add_argument("--cycle"); ap.add_argument("--round"); ap.add_argument("--glob", default="/mnt/user-uploads/*SAEB*.xlsx")
a = ap.parse_args()

def norm(s): return re.sub(r"\s+", " ", unicodedata.normalize("NFKD", str(s or "")).encode("ascii", "ignore").decode().lower()).strip(" .")
def q(v): return "NULL" if v is None else "'" + str(v).replace("'", "''") + "'"
def sh(cmd): return subprocess.run(["psql", "-Atc", cmd], capture_output=True, text=True, check=True).stdout.splitlines()

schools = {}
for l in sh("select distinct on (school_id) school_id, official_name from institutional_school_record_versions order by school_id, version_number desc"):
    sid, nm = l.split("|", 1); schools.setdefault(re.sub(r"^(e\.? ?m\.?|escola municipal) ", "", norm(nm)), set()).add(sid)
roster = {}
for l in sh("select e.school_id, s.id, s.display_name from school_enrollments e join institutional_students s on s.id = e.student_id"):
    sid, stid, nm = l.split("|", 2); roster.setdefault((sid, norm(nm)), set()).add(stid)

stats = dict(arquivos=0, abas=0, linhas=0, sem_dado=0, revisao_nome=0, sem_correspondencia=0, duplicado=0, escola_nao_identificada=0, ja_preparadas=0)
out = ["BEGIN;"]
for f in sorted(glob.glob(a.glob)):
    try:
        wbf = openpyxl.load_workbook(f, read_only=True); wbv = openpyxl.load_workbook(f, read_only=True, data_only=True)
    except Exception: continue
    sha = hashlib.sha256(open(f, "rb").read()).hexdigest(); name = f.split("/")[-1]; used = False
    for ws in wbv.worksheets:
        if "GRÁFICO" in ws.title.upper(): continue
        rows = list(ws.iter_rows(max_row=200, values_only=True))
        if len(rows) < 8 or "NOME" not in [str(x or "").strip() for x in rows[2]]: continue
        hdr = [str(x or "").strip() for x in rows[2]]
        c_name = hdr.index("NOME"); c_q = hdr.index("QUESTÕES") if "QUESTÕES" in hdr else c_name + 1
        c_tot = next((i for i, x in enumerate(hdr) if x.upper().startswith("TOTAL DE ACERTO")), None)
        qcols = list(range(c_q, c_tot if c_tot else len(hdr)))
        descr = [str(rows[4][i] or "") if i < len(rows[4]) else "" for i in qcols]
        school_label = next((str(x) for x in rows[0][2:] if x and not str(x).startswith("=")), None)
        cand = schools.get(re.sub(r"^(e\.? ?m\.?|escola municipal) ", "", norm(school_label))) if school_label else None
        school = next(iter(cand)) if cand and len(cand) == 1 else None
        if not school: stats["escola_nao_identificada"] += 1
        frow = list(wbf[ws.title].iter_rows(min_row=8, max_row=8, values_only=True))
        formula = str(frow[0][c_tot]) if frow and c_tot is not None and frow[0][c_tot] is not None else None
        class_label = str(rows[6][5]) if len(rows[6]) > 5 and rows[6][5] is not None else None
        stats["abas"] += 1; used = True; seen = set(); vals = []
        for n, r in enumerate(rows[7:], 8):
            nm = r[c_name] if c_name < len(r) else None
            if not nm or not str(nm).strip(): continue
            stats["linhas"] += 1
            ans = [None if i >= len(r) or r[i] in (None, "") else str(r[i]) for i in qcols]
            key = norm(nm); cands = sorted(roster.get((school, key), set())) if school else []
            if all(x is None for x in ans): st = "sem-dado"
            elif key in seen: st = "duplicado"
            elif cands: st = "revisao-nome"
            else: st = "sem-correspondencia"
            seen.add(key); stats[st.replace("-", "_")] += 1
            payload = {"answers": ans, "descriptors": descr, "present": r[3] if len(r) > 3 else None, "participates": r[2] if len(r) > 2 else None,
                       "total_local": r[c_tot] if c_tot is not None and c_tot < len(r) else None}
            vals.append(f"({n},{q(class_label)},{q(str(nm).strip())},NULL,NULL,{q(json.dumps(payload, ensure_ascii=False, default=str))}::jsonb,{q(st)},{q('{' + ','.join(cands) + '}')}::text[])")
        bid = f"(SELECT id FROM public.aa_import_batches WHERE adapter_key={q(AD)} AND adapter_version={VER} AND sha256={q(sha)} AND sheet={q(ws.title)})"
        out.append(f"INSERT INTO public.aa_import_batches(adapter_key,adapter_version,file_name,sha256,sheet,program_key,edition_year,cycle,round,school_label,measure_kind,formula_note,executor_label) VALUES ({q(AD)},{VER},{q(name)},{q(sha)},{q(ws.title)},{q(a.program)},{a.year},{q(a.cycle)},{q(a.round)},{q(school_label)},'acertos-locais',{q(formula)},{q(EX)}) ON CONFLICT DO NOTHING;")
        if vals:
            out.append("INSERT INTO public.aa_import_rows(batch_id,row_no,class_label,student_name,inep,enrollment_ref,payload,match_state,candidate_student_ids) SELECT " + bid + ", v.* FROM (VALUES " + ",".join(vals) + ") v(n,c,s,i,e,p,m,k) ON CONFLICT DO NOTHING;")
        out.append(f"INSERT INTO public.aa_import_batch_events(batch_id,seq,event,executor_label) SELECT {bid},1,'preparado',{q(EX)} ON CONFLICT DO NOTHING;")
    if used: stats["arquivos"] += 1
out.append("COMMIT;")
if a.sql: print("\n".join(out))
else: print(json.dumps(stats, ensure_ascii=False, indent=1))
