#!/usr/bin/env python3
"""Prévia (não grava) da atualização cadastral a partir dos "Banco de dados" das escolas.
Casa aluno só por INEP exato (identificador da pessoa já registrado); nome nunca casa.
Saída: contagens agregadas, sem nenhum dado pessoal."""
import glob, hashlib, json, re, subprocess, collections
import openpyxl

files = sorted(glob.glob("/mnt/user-uploads/Banco_de_dados_-_*.xlsx"))
inep_db = {}
out = subprocess.run(["psql", "-Atc", "select i.value, sp.student_id from institutional_person_identifiers i join institutional_student_persons sp on sp.person_id=i.person_id where i.identifier_kind='inep-pessoa'"], capture_output=True, text=True).stdout
for l in out.splitlines():
    v, s = l.split("|"); inep_db.setdefault(re.sub(r"\D", "", v), set()).add(s)

tot = collections.Counter(); per = []
for f in files:
    sha = hashlib.sha256(open(f, "rb").read()).hexdigest()
    wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
    if "ALUNOS" not in wb.sheetnames: per.append(dict(file=f.split("/")[-1], erro="sem aba ALUNOS")); continue
    rows = list(wb["ALUNOS"].iter_rows(values_only=True))
    h = [str(c or "").strip() for c in rows[0]]
    ix = {k: next((i for i, x in enumerate(h) if x.lower().startswith(k)), None) for k in ["nome do aluno", "identificação/inep", "cpf", "data de nascimento", "sexo", "cor/raça", "nee", "situação da matr"]}
    c = collections.Counter(); seen = collections.Counter()
    for r in rows[1:]:
        r = tuple(r) + (None,) * (len(h) - len(r))
        if ix["nome do aluno"] is None or not r[ix["nome do aluno"]]: continue
        c["linhas"] += 1
        g = lambda k: (r[ix[k]] if ix[k] is not None and ix[k] < len(r) else None)
        inep = re.sub(r"\D", "", str(g("identificação/inep") or ""))
        if not inep: c["sem_inep"] += 1; continue
        seen[inep] += 1
        m = inep_db.get(inep)
        if not m: c["inep_nao_encontrado"] += 1
        elif len(m) > 1: c["inep_ambiguo"] += 1
        else: c["casado"] += 1
        if g("data de nascimento"): c["com_nascimento"] += 1
        if str(g("sexo") or "").strip().upper() in ("M", "F"): c["com_sexo"] += 1
        if g("cpf"): c["com_cpf"] += 1
    c["inep_repetido_no_arquivo"] = sum(1 for v in seen.values() if v > 1)
    tot.update(c); per.append(dict(file=f.split("/")[-1], sha256=sha[:12], **c))
print(json.dumps(dict(arquivos=len(files), total=tot, por_arquivo=per), ensure_ascii=False, indent=1))
