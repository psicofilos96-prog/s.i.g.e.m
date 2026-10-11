#!/usr/bin/env python3
"""Carga técnica reexecutável das planilhas "Banco de dados" das escolas como OBSERVAÇÃO cadastral.
Só linhas cujo INEP casa exatamente com UMA pessoa-aluno do SIGEM; nome nunca casa. CPF não é gravado.
Uso: python3 scripts/school-db-import.py > /tmp/sro.sql && psql -v ON_ERROR_STOP=1 -f /tmp/sro.sql"""
import datetime, glob, hashlib, re, subprocess, sys
import openpyxl

EX = "importacao-tecnica:scripts/school-db-import.py"
inep_db = {}
for l in subprocess.run(["psql", "-Atc", "select i.value, sp.student_id from institutional_person_identifiers i join institutional_student_persons sp on sp.person_id=i.person_id where i.identifier_kind='inep-pessoa'"], capture_output=True, text=True, check=True).stdout.splitlines():
    v, s = l.split("|"); inep_db.setdefault(re.sub(r"\D", "", v), set()).add(s)

def q(v):
    if v is None: return "NULL"
    return "'" + str(v).replace("'", "''") + "'"
def txt(v):
    if v is None: return None
    s = re.sub(r"\s+", " ", str(v)).strip()
    return None if s in ("", "-", "X", "x", "None") else s
def date(v):
    if isinstance(v, datetime.datetime): return v.date().isoformat()
    if isinstance(v, datetime.date): return v.isoformat()
    m = re.fullmatch(r"(\d{1,2})/(\d{1,2})/(\d{4})", str(v or "").strip())
    return f"{m[3]}-{int(m[2]):02d}-{int(m[1]):02d}" if m else None

stats = dict(linhas=0, gravaveis=0, sem_inep=0, nao_encontrado=0, ambiguo=0)
vals = []
for f in sorted(glob.glob("/mnt/user-uploads/Banco_de_dados_-_*.xlsx")):
    sha = hashlib.sha256(open(f, "rb").read()).hexdigest(); name = f.split("/")[-1]
    wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
    if "ALUNOS" not in wb.sheetnames: continue
    rows = list(wb["ALUNOS"].iter_rows(values_only=True)); h = [str(c or "").strip().lower() for c in rows[0]]
    col = lambda k: next((i for i, x in enumerate(h) if x.startswith(k)), None)
    C = {k: col(k) for k in ["nome do aluno", "identificação/inep", "data de nascimento", "sexo", "cor/raça", "nee", "pai", "mãe", "naturalidade", "uf naturalidade", "nacionalidade", "responsável", "profissão do resp", "telefone 1", "telefone 2", "telefone 3", "situação da matr", "turma", "ano de escolaridade"]}
    for n, r in enumerate(rows[1:], 2):
        r = tuple(r) + (None,) * (len(h) - len(r)); g = lambda k: r[C[k]] if C[k] is not None else None
        if not txt(g("nome do aluno")): continue
        stats["linhas"] += 1
        inep = re.sub(r"\D", "", str(g("identificação/inep") or ""))
        if not inep: stats["sem_inep"] += 1; continue
        m = inep_db.get(inep)
        if not m: stats["nao_encontrado"] += 1; continue
        if len(m) > 1: stats["ambiguo"] += 1; continue
        stats["gravaveis"] += 1
        phones = [p for p in (txt(g(k)) for k in ["telefone 1", "telefone 2", "telefone 3"]) if p]
        arr = "ARRAY[" + ",".join(q(p) for p in phones) + "]::text[]" if phones else "'{}'::text[]"
        vals.append("(" + ",".join([q(next(iter(m))), q(name), q(sha), "'ALUNOS'", str(n), "'inep-pessoa-exato'", q(EX), q(txt(g("nome do aluno"))), q(date(g("data de nascimento"))) + "::date",
            q(txt(g("sexo"))), q(txt(g("cor/raça"))), q(txt(g("nee"))), q(txt(g("pai"))), q(txt(g("mãe"))), q(txt(g("naturalidade"))), q(txt(g("uf naturalidade"))),
            q(txt(g("nacionalidade"))), q(txt(g("responsável"))), q(txt(g("profissão do resp"))), arr, q(txt(g("situação da matr"))), q(txt(g("turma"))), q(txt(g("ano de escolaridade")))]) + ")")

cols = "student_id,source_file,source_sha256,source_sheet,source_row,matched_by,executor_label,civil_name,birth_date,sex_literal,race_literal,nee_literal,father_name,mother_name,birthplace,birthplace_uf,nationality,guardian_name,guardian_occupation,phones,enrollment_status_literal,class_literal,grade_literal"
for i in range(0, len(vals), 500):
    print(f"INSERT INTO public.student_registration_observations({cols}) VALUES {','.join(vals[i:i+500])} ON CONFLICT (source_sha256, source_sheet, source_row) DO NOTHING;")
print("-- " + str(stats), file=sys.stderr)
